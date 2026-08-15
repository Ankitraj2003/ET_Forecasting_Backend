from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def create_forecast_excel(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ETo Forecast"
    title_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="D9EAF7")

    ws["A1"] = "ETo Forecast — Hargreaves-Samani + FAO-56 Penman-Monteith"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws.merge_cells("A1:P1")

    metadata = [("Location", data["resolved_location"]), ("Country", data.get("country") or ""),
                ("Latitude", data["latitude"]), ("Longitude", data["longitude"]),
                ("Timezone", data.get("timezone") or ""), ("Selected Method", data["method"]),
                ("Weather Source", data["weather_source"])]
    row = 3
    for label, value in metadata:
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, value)
        row += 1

    table_row = row + 1
    headers = ["Date", "DOY", "Tmax °C", "Tmin °C", "Tmean °C", "RHmax %", "RHmin %",
               "Wind 10m m/s", "Rs MJ/m²/day", "Pressure kPa", "Ra MJ/m²/day",
               "HS ETo mm/day", "FAO-56 PM ETo mm/day", "Selected ETo mm/day",
               "Rn MJ/m²/day", "VPD kPa"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(table_row, c, h); cell.font = Font(bold=True); cell.fill = header_fill; cell.alignment = Alignment(horizontal="center")
    for r, item in enumerate(data["forecast"], table_row + 1):
        values = [item["date"], item["day_of_year"], item["tmax_c"], item["tmin_c"], item["tmean_c"],
                  item["rh_max_pct"], item["rh_min_pct"], item["wind_speed_10m_m_s"], item["solar_radiation_mj_m2_day"],
                  item["pressure_kpa"], item["ra_mj_m2_day"], item["hargreaves_samani_eto_mm_day"],
                  item["penman_monteith_eto_mm_day"], item["eto_mm_day"], item["pm_net_radiation_mj_m2_day"], item["pm_vapor_pressure_deficit_kpa"]]
        for c, value in enumerate(values, 1): ws.cell(r, c, value)

    method = wb.create_sheet("Methods")
    method["A1"] = "ETo Methods Used"; method["A1"].font = Font(bold=True, size=14)
    method["A3"] = "Hargreaves-Samani"; method["A3"].font = Font(bold=True)
    method["A4"] = "ETo = 0.0023 × Ra × (Tmean + 17.8) × √(Tmax − Tmin)"
    method["A6"] = "FAO-56 Penman-Monteith"; method["A6"].font = Font(bold=True)
    method["A7"] = "ETo = [0.408Δ(Rn−G) + γ(900/(T+273))u₂(es−ea)] / [Δ + γ(1+0.34u₂)]"
    method["A9"] = "PM inputs"; method["B9"] = "Tmax/Tmin, RHmax/RHmin, wind speed, solar radiation, pressure, latitude and day of year."
    method["A11"] = "Important"; method["B11"] = "Both methods are calculated for every day. The selected method controls the primary ETo column."
    method["A13"] = "Reference"; method["B13"] = "FAO-56 reference evapotranspiration methodology (grass reference surface)."

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cells in sheet.columns:
            letter = get_column_letter(cells[0].column)
            width = max([len(str(c.value)) for c in cells if c.value is not None] + [12])
            sheet.column_dimensions[letter].width = min(width + 2, 45)
    output = BytesIO(); wb.save(output); return output.getvalue()


def create_etc_forecast_excel(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active; ws.title = "ETc Forecast"
    ws["A1"] = "7-Day Crop Evapotranspiration Forecast (ETc)"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells("A1:S1")
    meta=[("Location",data["resolved_location"]),("Crop",data["crop"]),("Sowing Date",data["sowing_date"]),
          ("Forecast Start",data["forecast_start_date"]),("ETo Method",data["method"]),
          ("Kc Source",data["kc_source"]),("Kc Source URL",data.get("kc_source_url", ""))]
    rr=3
    for k,v in meta:
        ws.cell(rr,1,k).font=Font(bold=True); ws.cell(rr,2,v); rr+=1
    r0=rr+1
    headers=["Date","Crop Day","Growth Stage","Kc","ETo mm/day","ETc mm/day","Tmax °C","Tmin °C","Tmean °C","RHmax %","RHmin %","Wind 10m m/s","Solar Radiation MJ/m²/day","Ra MJ/m²/day","HS ETo mm/day","FAO-56 PM ETo mm/day","Selected Method","Latitude","Longitude"]
    for c,h in enumerate(headers,1):
        x=ws.cell(r0,c,h); x.font=Font(bold=True); x.fill=PatternFill("solid", fgColor="D9EAF7"); x.alignment=Alignment(horizontal="center")
    for r,item in enumerate(data["forecast"],r0+1):
        vals=[item["date"],item["crop_day"],item["growth_stage"],item["kc"],item["eto_mm_day"],item["etc_mm_day"],item["tmax_c"],item["tmin_c"],item["tmean_c"],item["rh_max_pct"],item["rh_min_pct"],item["wind_speed_10m_m_s"],item["solar_radiation_mj_m2_day"],item["ra_mj_m2_day"],item["hargreaves_samani_eto_mm_day"],item["penman_monteith_eto_mm_day"],item["selected_method"],data["latitude"],data["longitude"]]
        for c,v in enumerate(vals,1): ws.cell(r,c,v)
    note=wb.create_sheet("Kc Reference")
    note["A1"]="FAO-56 Kc reference"; note["A1"].font=Font(bold=True,size=14)
    note["A3"]="Formula"; note["B3"]="ETc = Kc × ETo"
    note["A4"]="Kc source"; note["B4"]=data["kc_source"]
    note["A5"]="Source URL"; note["B5"]=data.get("kc_source_url","")
    note["A7"]="Important"; note["B7"]="FAO-56 Kc values are reference values, not live weather data. Daily Kc is constructed from the segmented crop-coefficient curve using the crop profile and stage lengths. Local phenology and climate adjustments may be needed for field-level accuracy."
    for sh in wb.worksheets:
        sh.freeze_panes="A2"
        for cells in sh.columns:
            letter=get_column_letter(cells[0].column); width=max([len(str(c.value)) for c in cells if c.value is not None]+[12]); sh.column_dimensions[letter].width=min(width+2,45)
    out=BytesIO(); wb.save(out); return out.getvalue()
